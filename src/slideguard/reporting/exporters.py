from __future__ import annotations

from collections import Counter
from html import escape
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Font

from slideguard.preview.svg_builder import PreviewObject, build_svg
from slideguard.scan.models import ScanResult


def export_html(result: ScanResult, destination: Path, *, software_version: str = "0.1.0") -> None:
    _validate_destination(destination, ".html")
    severity = Counter(found.severity.value for found in result.issues)
    rules = Counter(found.rule_id for found in result.issues)
    affected_pages = len({found.slide_index for found in result.issues if found.slide_index > 0})
    fixable = sum(found.can_auto_fix for found in result.issues)
    rows = "".join(_html_issue_row(result, found) for found in result.issues)
    rule_summary = "".join(
        f"<li>{escape(rule_id)}：{count}</li>" for rule_id, count in sorted(rules.items())
    )
    rule_options = "".join(
        f"<option value='{escape(rule_id, quote=True)}'>{escape(rule_id)}</option>"
        for rule_id in sorted(rules)
    )
    incomplete = "<div class='incomplete'>扫描未完成</div>" if not result.complete else ""
    execution = _html_execution_summary(result)
    comparison = _html_repair_comparison(result)
    lexicon_warning = (
        "<div class='incomplete'>敏感词库为空，本项未发现问题不代表无敏感内容。</div>"
        if result.sensitive_lexicon_empty else ""
    )
    unsupported_warning = _html_unsupported_summary(result)
    document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>SlideGuard 质检报告</title>
<style>body{{font-family:'Microsoft YaHei',sans-serif;margin:32px;color:#222}}h1{{color:#c00000}}.incomplete{{padding:16px;background:#ffd9d9;color:#8b0000;font-weight:bold}}.filters{{display:flex;gap:8px;margin:18px 0}}.filters input,.filters select{{padding:7px;border:1px solid #aaa}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ccc;padding:8px;text-align:left;vertical-align:top}}th{{background:#f2f2f2}}.S1{{color:#a00000;font-weight:bold}}.preview{{width:240px;max-height:140px}}</style></head>
<body><h1>SlideGuard 质检报告</h1>{incomplete}{lexicon_warning}{unsupported_warning}
<dl><dt>文件名</dt><dd>{escape(result.snapshot.file_identity.path.name)}</dd><dt>路径</dt><dd>{escape(str(result.snapshot.file_identity.path))}</dd><dt>大小</dt><dd>{result.snapshot.file_identity.size_bytes} bytes</dd><dt>页数</dt><dd>{len(result.snapshot.slides)}</dd><dt>扫描模式</dt><dd>{escape(result.mode.value)}</dd><dt>完整性</dt><dd>{'完整' if result.complete else '未完成'}</dd><dt>规则集</dt><dd>{escape(result.rule_set_version)}</dd><dt>软件版本</dt><dd>{escape(software_version)}</dd><dt>开始时间</dt><dd>{result.started_at.isoformat()}</dd><dt>结束时间</dt><dd>{result.finished_at.isoformat()}</dd></dl>
{execution}<p>S1 {severity['S1']} / S2 {severity['S2']} / S3 {severity['S3']} / S4 {severity['S4']}；涉及页面 {affected_pages}；可自动修复 {fixable}</p>{comparison}<ul>{rule_summary}</ul>
<div class="filters"><input id="report-search" type="search" placeholder="搜索问题明细"><select id="report-severity"><option value="">全部级别</option><option>S1</option><option>S2</option><option>S3</option><option>S4</option></select><select id="report-rule"><option value="">全部规则</option>{rule_options}</select><span id="visible-count"></span></div>
<table><thead><tr><th>级别</th><th>规则</th><th>页码</th><th>状态</th><th>修复后新增</th><th>实际值</th><th>标准值</th><th>依据</th><th>建议</th><th>预览</th></tr></thead><tbody>{rows}</tbody></table>{_REPORT_SCRIPT}</body></html>"""
    _atomic_text(destination, document)


def export_xlsx(result: ScanResult, destination: Path, *, software_version: str = "0.1.0") -> None:
    _validate_destination(destination, ".xlsx")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "扫描摘要"
    summary.append(("项目", "值"))
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    severity = Counter(found.severity.value for found in result.issues)
    rules = Counter(found.rule_id for found in result.issues)
    affected_pages = len({found.slide_index for found in result.issues if found.slide_index > 0})
    fixable = sum(found.can_auto_fix for found in result.issues)
    values = (
        ("文件名", result.snapshot.file_identity.path.name),
        ("路径", str(result.snapshot.file_identity.path)),
        ("大小(bytes)", result.snapshot.file_identity.size_bytes),
        ("页数", len(result.snapshot.slides)),
        ("扫描模式", result.mode.value),
        ("完整性", "完整" if result.complete else "扫描未完成"),
        ("规则集版本", result.rule_set_version),
        ("软件版本", software_version),
        ("开始时间", result.started_at.isoformat()),
        ("结束时间", result.finished_at.isoformat()),
        ("敏感词库状态", "为空：未发现问题不代表无敏感内容" if result.sensitive_lexicon_empty else "已配置或未执行R010"),
        ("未支持对象总数", len(result.snapshot.unsupported_objects)),
        *_xlsx_unsupported_summary(result),
        ("解析失败范围总数", len(result.snapshot.parse_failures)),
        ("解析失败页", "、".join(str(page) for page in sorted({item.slide_index for item in result.snapshot.parse_failures})) or "无"),
        ("请求执行规则", "、".join(result.requested_rules)),
        ("已完成规则", "、".join(result.completed_rules)),
        ("失败规则", "；".join(f"{item.rule_id}: {item.message}" for item in result.failures) or "无"),
        ("涉及问题页面数", affected_pages),
        ("可自动修复问题数", fixable),
        *( (level, severity[level]) for level in ("S1", "S2", "S3", "S4") ),
        *( (f"{rule_id} 问题数", count) for rule_id, count in sorted(rules.items()) ),
        *_xlsx_repair_comparison(result),
    )
    for row in values:
        summary.append(row)
    details = workbook.create_sheet("问题清单")
    headers = ("问题ID", "级别", "规则", "页码", "状态", "可自动修复", "实际值", "标准值", "标准来源", "判断依据", "修改建议", "对象ID", "修复后新增")
    details.append(headers)
    for cell in details[1]:
        cell.font = Font(bold=True)
    for found in result.issues:
        details.append(tuple(_safe_excel(value) for value in (
            found.issue_id, found.severity.value, found.rule_id, found.slide_index,
            found.status.value, "是" if found.can_auto_fix else "否", found.actual_value,
            found.expected_value, found.standard_source, found.evidence, found.suggestion,
            ", ".join(found.object_keys), "是" if found.introduced_by_repair else "否",
        )))
    for sheet in (summary, details):
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    destination.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        temporary_path.rename(destination)
    finally:
        temporary_path.unlink(missing_ok=True)


def default_report_name(result: ScanResult, extension: str) -> str:
    timestamp = result.finished_at.astimezone().strftime("%Y%m%d_%H%M%S")
    return f"{result.snapshot.file_identity.path.stem}_SlideGuard_{timestamp}.{extension.lstrip('.')}"


def _html_repair_comparison(result: ScanResult) -> str:
    comparison = result.repair_comparison
    if comparison is None:
        return ""
    return (
        "<p><strong>修复前后对比：</strong>"
        f"选中 {comparison.selected_count}；已修复 {comparison.fixed_count}；"
        f"未修复 {comparison.unresolved_count}；修复后新增 {comparison.introduced_count}</p>"
    )


def _html_execution_summary(result: ScanResult) -> str:
    requested = escape("、".join(result.requested_rules))
    completed = escape("、".join(result.completed_rules))
    failures = "".join(
        f"<li>{escape(item.rule_id)}：{escape(item.message)}</li>"
        for item in result.failures
    )
    failure_block = f"<ul>{failures}</ul>" if failures else "<span>无</span>"
    return (
        f"<p><strong>请求执行规则：</strong>{requested}</p>"
        f"<p><strong>已完成规则：</strong>{completed}</p>"
        f"<div><strong>失败规则：</strong>{failure_block}</div>"
    )


def _unsupported_counts(result: ScanResult) -> Counter[tuple[int, str]]:
    return Counter(
        (item.slide_index, item.object_type)
        for item in result.snapshot.unsupported_objects
    )


def _html_unsupported_summary(result: ScanResult) -> str:
    counts = _unsupported_counts(result)
    if not counts:
        return ""
    items = "".join(
        f"<li>第 {page} 页 {escape(object_type)}：{count} 个</li>"
        for (page, object_type), count in sorted(counts.items())
    )
    return (
        "<div class='incomplete'><strong>存在未支持检查的对象，以下范围不得视为检查通过：</strong>"
        f"<ul>{items}</ul></div>"
    )


def _xlsx_unsupported_summary(result: ScanResult) -> tuple[tuple[str, object], ...]:
    return tuple(
        (f"未支持对象-第{page}页-{object_type}", count)
        for (page, object_type), count in sorted(_unsupported_counts(result).items())
    )


def _xlsx_repair_comparison(result: ScanResult) -> tuple[tuple[str, int], ...]:
    comparison = result.repair_comparison
    if comparison is None:
        return ()
    return (
        ("修复选中问题数", comparison.selected_count),
        ("已修复问题数", comparison.fixed_count),
        ("未修复问题数", comparison.unresolved_count),
        ("修复后新增问题数", comparison.introduced_count),
    )


def _html_issue_row(result: ScanResult, found) -> str:  # type: ignore[no-untyped-def]
    attributes = (
        f" data-severity='{escape(found.severity.value, quote=True)}'"
        f" data-rule='{escape(found.rule_id, quote=True)}'"
    )
    return f"<tr{attributes}>" + "".join(
        f"<td>{escape(str(value))}</td>" for value in (
            found.severity.value, found.rule_id, found.slide_index, found.status.value,
            "是" if found.introduced_by_repair else "否", found.actual_value,
            found.expected_value, found.evidence, found.suggestion,
        )
    ) + f"<td>{_slide_svg(result, found)}</td></tr>"


_REPORT_SCRIPT = """<script>
(() => {
  const search = document.querySelector('#report-search');
  const severity = document.querySelector('#report-severity');
  const rule = document.querySelector('#report-rule');
  const rows = [...document.querySelectorAll('tbody tr')];
  const count = document.querySelector('#visible-count');
  const apply = () => {
    const query = search.value.trim().toLowerCase();
    let visible = 0;
    rows.forEach((row) => {
      const show = (!query || row.textContent.toLowerCase().includes(query))
        && (!severity.value || row.dataset.severity === severity.value)
        && (!rule.value || row.dataset.rule === rule.value);
      row.hidden = !show;
      if (show) visible += 1;
    });
    count.textContent = `显示 ${visible}/${rows.length} 个问题`;
  };
  [search, severity, rule].forEach((control) => control.addEventListener('input', apply));
  apply();
})();
</script>"""


def _slide_svg(result: ScanResult, found) -> str:  # type: ignore[no-untyped-def]
    if found.slide_index < 1 or found.slide_index > len(result.snapshot.slides):
        return "预览不可用"
    slide = result.snapshot.slides[found.slide_index - 1]
    objects = tuple(_preview_objects(slide.objects))
    return build_svg(
        slide_width_pt=result.snapshot.slide_width_pt,
        slide_height_pt=result.snapshot.slide_height_pt,
        objects=objects,
        highlighted_ids=frozenset(found.object_keys[:1]),
        reference_ids=frozenset(found.object_keys[1:] if found.rule_id == "R007" else ()),
        page_highlight=not found.object_keys,
    ).replace("<svg ", "<svg class='preview' ", 1)


def _preview_objects(objects):  # type: ignore[no-untyped-def]
    for obj in objects:
        yield PreviewObject(
            obj.key,
            obj.bounds_pt.left,
            obj.bounds_pt.top,
            obj.bounds_pt.width,
            obj.bounds_pt.height,
            obj.text_frame.text if obj.text_frame is not None else "",
        )
        yield from _preview_objects(obj.children)


def _safe_excel(value):  # type: ignore[no-untyped-def]
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _validate_destination(destination: Path, suffix: str) -> None:
    if destination.suffix.lower() != suffix:
        raise ValueError(f"报告扩展名必须为 {suffix}")
    if destination.exists():
        raise FileExistsError("目标报告已存在，不允许覆盖")


def _atomic_text(destination: Path, content: str) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile("w", encoding="utf-8", prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent, delete=False) as temporary:
        temporary.write(content)
        temporary_path = Path(temporary.name)
    try:
        temporary_path.rename(destination)
    finally:
        temporary_path.unlink(missing_ok=True)
